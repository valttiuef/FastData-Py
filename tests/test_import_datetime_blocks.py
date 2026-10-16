"""Regression coverage for separate clocks and repeated Excel measurements."""
# @ai(gpt-6, codex, regression-tests, 2026-10-16)
import datetime
import os
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path.cwd() / "src"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from backend.models import HeaderRoles, ImportOptions
from backend.importing.parsers import _build_tall_from_raw, parse_file_to_tall, read_raw_excel
from backend.importing.utils import _normalize_time_series
from backend.data_db.database import Database
from frontend.tabs.data.import_preview_logic import build_import_preview_payload


@pytest.mark.parametrize("clock", [datetime.time(10, 20), "10:20", "10.20", 10.2, 31 / 72])
def test_separate_clock_formats(clock):
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    raw = pd.DataFrame([["Date", "Time", "Value"], ["02/01/2025", clock, 5.25]])
    tall, headers = _build_tall_from_raw(raw, ImportOptions(excel_header_rows=1, header_roles=HeaderRoles(0, None, None, None)))
    assert headers == 1
    assert tall.ts.tolist() == [pd.Timestamp("2025-01-02 10:20")]
    assert tall.value.tolist() == [5.25]
    assert tall.base_name.tolist() == ["Value"]


def test_continuation_rows_preserve_samples_and_reset_clock_at_new_date():
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    raw = pd.DataFrame([
        ["Date", "Time", "Value"],
        ["02/01/2025", "10:20", 5.25],
        [None, None, 6.25], [None, None, 7.25],
        ["03/01/2025", None, 999],
        ["03/01/2025", "17:45", 8.25],
    ])
    opts = ImportOptions(excel_header_rows=1, fill_down_timestamps=True, average_repeated_measurements=False, header_roles=HeaderRoles(0, None, None, None))
    tall, _ = _build_tall_from_raw(raw, opts)
    assert tall.value.tolist() == [5.25, 6.25, 7.25, 999, 8.25]
    assert tall.type.tolist() == ["Sample 1", "Sample 2", "Sample 3", "Sample 1", "Sample 1"]
    assert tall.ts.tolist()[:3] == [pd.Timestamp("2025-01-02 10:20")] * 3
    assert tall.ts.iloc[-1] == pd.Timestamp("2025-01-03 17:45")
    assert tall.ts.iloc[-2] == pd.Timestamp("2025-01-03")
    assert tall.attrs["import_warnings"]


def test_clock_validation_and_dot_toggle():
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    assert _normalize_time_series(pd.Series(["24:00", "10:60", "10.20"]), dot_time_as_colon=False).isna().all()


def test_merged_workbook_headers_and_repeated_values(tmp_path):
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, None, None, "Thickness", None])
    sheet.append(["Date", "Time", "Operator", "Left", "Right"])
    sheet.append([datetime.datetime(2025, 1, 2), datetime.time(10, 20), "Operator A", 1.5, 1.6])
    sheet.append([None, None, None, 1.7, 1.8])
    for cell_range in ("D1:E1", "A3:A4", "B3:B4", "C3:C4"):
        sheet.merge_cells(cell_range)
    path = tmp_path / "merged.xlsx"
    workbook.save(path)
    workbook.close()
    raw = next(iter(read_raw_excel(path).values()))
    from PySide6.QtWidgets import QApplication
    from frontend.windows.import_options_dialog import ImportOptionsDialog

    app = QApplication.instance() or QApplication([])
    payload = build_import_preview_payload(file_path=str(path), csv_delimiter=None, csv_decimal=None, csv_encoding=None, base_header_index=None)
    dialog = ImportOptionsDialog([path], database_model=None, data_view_model=None)
    dialog._apply_preview_payload(payload, dialog._preview_request_id)
    opts = dialog.build_options()
    assert opts.header_roles.unit_row is None
    assert opts.header_roles.type_row is None
    assert opts.average_repeated_measurements is True
    dialog.average_repeats_cb.setChecked(False)
    assert dialog.build_options().average_repeated_measurements is False
    dialog.close()
    tall, _ = _build_tall_from_raw(raw, opts)
    assert tall.value.tolist() == pytest.approx([1.6, 1.7])
    assert set(tall.ts) == {pd.Timestamp("2025-01-02 10:20")}
    assert set(tall.base_name) == {"Thickness"}
    assert set(tall.source) == {"Left", "Right"}
    assert set(tall.Operator) == {"Operator A"}
    assert tall.type.isna().all()


def test_preview_detects_unmerged_continuation_rows(tmp_path):
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    path = tmp_path / "continuations.csv"
    path.write_text("Date,Time,Value\n02/01/2025,10.20,5.25\n,,6.25\n,,7.25\n", encoding="utf-8")
    payload = build_import_preview_payload(file_path=str(path), csv_delimiter=None, csv_decimal=None, csv_encoding=None, base_header_index=None)
    assert payload["error"] is None
    assert payload["fill_down_timestamps"] is True
    assert payload["date_column"] == "Date"
    assert payload["time_column"] == "Time"


@pytest.mark.parametrize("use_duckdb", [True, False])
def test_separate_clock_csv_survives_database_import(tmp_path, use_duckdb):
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    path = tmp_path / "clocks.csv"
    path.write_text("Date,Time,Value\n2025-01-02,10:20,5.25\n2025-01-02,17:45,8.25\n", encoding="utf-8")
    db = Database(tmp_path / "data.duckdb")
    try:
        db.import_file(path, ImportOptions(use_duckdb_csv_import=use_duckdb, header_roles=HeaderRoles(0, None, None, None)))
        rows = db.con.execute("SELECT ts, value FROM measurements ORDER BY ts").fetchall()
        assert rows == [(datetime.datetime(2025, 1, 2, 10, 20), 5.25), (datetime.datetime(2025, 1, 2, 17, 45), 8.25)]
    finally:
        db.close()


@pytest.mark.parametrize("average", [True, False])
@pytest.mark.parametrize("header_count,date_as_text", [(5, False), (4, True)])
def test_generated_excel_preview_and_database_preserve_every_measurement(tmp_path, header_count, date_as_text, average):
    # @ai(gpt-6, codex, integration-test, 2026-10-16)
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "Measurement log")
    sheet.cell(header_count - 1, 4, "Length")
    sheet.merge_cells(start_row=header_count - 1, start_column=4, end_row=header_count - 1, end_column=5)
    for col, label in enumerate(["Date", "Time", "Operator", "Left", "Right"], 1):
        sheet.cell(header_count, col, label)
    for block, clock in enumerate([datetime.time(9, 15), datetime.time(17, 45)]):
        first_row = header_count + 1 + block * 3
        date = "14.05.2030" if date_as_text else datetime.datetime(2030, 5, 14)
        for col, value in enumerate([date, clock, "Operator A"], 1):
            sheet.cell(first_row, col, value)
            sheet.merge_cells(start_row=first_row, start_column=col, end_row=first_row + 2, end_column=col)
        for offset in range(3):
            sheet.cell(first_row + offset, 4, 10 + block * 3 + offset)
            sheet.cell(first_row + offset, 5, 20 + block * 3 + offset)
    path = tmp_path / "measurement_log.xlsx"
    workbook.save(path)
    workbook.close()
    payload = build_import_preview_payload(file_path=str(path), csv_delimiter=None, csv_decimal=None, csv_encoding=None, base_header_index=None)
    assert payload["error"] is None
    assert payload["header_rows"] == header_count
    assert payload["date_column"] == "Date"
    assert payload["time_column"] == "Time"
    assert payload["force_meta_columns_guess"] == ["Operator"]
    opts = ImportOptions(
        average_repeated_measurements=average,
        excel_header_rows=header_count, date_column=payload["date_column"], time_column=payload["time_column"],
        fill_down_timestamps=payload["fill_down_timestamps"], assume_dayfirst=payload["assume_dayfirst"],
        datetime_formats=payload["datetime_formats"], force_meta_columns=payload["force_meta_columns_guess"],
        header_roles=HeaderRoles(payload["base_header_index"], payload["source_header_index"], None, None),
    )
    tall = parse_file_to_tall(path, opts)[0][1]
    assert len(tall) == (4 if average else 12)
    first = tall[tall.ts == pd.Timestamp("2030-05-14 09:15")]
    assert len(first) == (2 if average else 6)
    if average:
        assert first.type.isna().all()
        assert sorted(first.value) == [11.0, 21.0]
    else:
        assert set(first.type) == {"Sample 1", "Sample 2", "Sample 3"}
    assert set(first.Operator) == {"Operator A"}
    assert set(first.base_name) == {"Length"}
    db = Database(tmp_path / "data.duckdb")
    try:
        db.import_file(path, opts)
        stored = db.con.execute("SELECT value FROM measurements ORDER BY value").fetchall()
        values = [11, 14, 21, 24] if average else [10, 11, 12, 13, 14, 15, 20, 21, 22, 23, 24, 25]
        assert stored == [(float(value),) for value in values]
    finally:
        db.close()


@pytest.mark.parametrize("fill_down", [True, False])
def test_average_repeated_rows_preserves_columns_types_and_missing_values(fill_down):
    # @ai(gpt-6, codex, regression-test, 2026-10-16)
    raw = pd.DataFrame([
        ["Date", "Temperature", "Temperature", "Operator"],
        [None, "Indoor", "Outdoor", None],
        [None, "C", "C", None],
        [None, "Measured", "Measured", None],
        ["02/01/2025", 2, 10, "A"],
        ["02/01/2025", 4, None, None],
        ["02/01/2025", None, 20, "B"],
        ["03/01/2025", 8, 30, "C"],
    ])
    opts = ImportOptions(
        excel_header_rows=4, header_roles=HeaderRoles(0, 1, 2, 3),
        fill_down_timestamps=fill_down, force_meta_columns=["Operator"],
    )
    tall, _ = _build_tall_from_raw(raw, opts)
    assert tall.value.tolist() == [3.0, 8.0, 15.0, 30.0]
    assert tall.source.tolist() == ["Indoor", "Indoor", "Outdoor", "Outdoor"]
    assert set(tall.unit) == {"C"}
    assert set(tall.type) == {"Measured"}
    assert tall.Operator.tolist() == ["A", "C", "A", "C"]
    assert tall.attrs["forced_meta_kinds"] == ["Operator"]
